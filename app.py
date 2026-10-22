import streamlit as st
import os
import io
import re
import html
import time
import random
import tempfile
from datetime import datetime, timezone, timedelta
from google import genai
from google.genai import errors as genai_errors
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

st.set_page_config(page_title="光文社ビジネスクリエーション部", layout="wide")

JST = timezone(timedelta(hours=9))

# =========================================================
# 503対策の設定
# =========================================================
# 上から順に試す。存在しないモデル(404)は自動で飛ばす。
MODEL_CANDIDATES = [
    "gemini-3.6-flash",          # メイン（社内で利用可能と確認済み）
    "gemini-3.6-flash-preview",  # 正式名に -preview が付く場合の保険
    "gemini-3.5-flash",          # 以下、混雑時の代替（存在しなければ自動で飛ばす）
    "gemini-3.1-flash-lite",
    "gemini-2.5-flash",
]
RETRYABLE_CODES = {429, 500, 503, 504}   # 混雑・一時障害とみなして再試行するコード
MAX_ATTEMPTS_PER_MODEL = 5               # 1モデルあたりの試行回数
BASE_WAIT_SEC = 10                       # 初回の待ち時間（以後倍々）
MAX_WAIT_SEC = 120                       # 1回の待ち時間の上限
FILE_PROCESS_TIMEOUT_SEC = 900           # 音声のサーバー側処理待ちの上限


# =========================================================
# 見た目（CSS）
# =========================================================
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Noto+Sans+JP:wght@400;500;700&display=swap');

html, body, [class*="css"], .stMarkdown, .stText, button, input, textarea {
    font-family: 'Noto Sans JP', 'Hiragino Kaku Gothic ProN', 'Yu Gothic', sans-serif;
}
.stApp { background: #F5F6F8; }
header[data-testid="stHeader"] { background: transparent; }
footer { visibility: hidden; }
.block-container { padding-top: 2.2rem; padding-bottom: 4rem; }

/* ヘッダー */
.bc-header {
    background: #1E293B;
    border-radius: 6px;
    padding: 30px 34px 28px;
    margin-bottom: 28px;
    border-top: 3px solid #B08D57;
}
.bc-kicker {
    color: #B08D57;
    font-size: 11px;
    letter-spacing: 0.22em;
    font-weight: 500;
    margin: 0 0 10px;
}
.bc-title {
    color: #F8FAFC;
    font-size: 34px;
    font-weight: 700;
    letter-spacing: 0.02em;
    margin: 0;
    padding: 0;
    line-height: 1.25;
}
.bc-desc {
    color: #94A3B8;
    font-size: 13.5px;
    line-height: 1.7;
    margin: 12px 0 0;
}

/* セクション見出し */
.bc-label {
    font-size: 12px;
    font-weight: 700;
    color: #475569;
    letter-spacing: 0.12em;
    margin: 6px 0 8px;
    padding-bottom: 6px;
    border-bottom: 1px solid #E2E8F0;
}

/* アップロード欄 */
[data-testid="stFileUploaderDropzone"] {
    background: #FFFFFF;
    border: 1.5px dashed #CBD5E1;
    border-radius: 6px;
}
[data-testid="stFileUploaderDropzone"]:hover { border-color: #1E293B; }

.bc-file {
    background: #FFFFFF;
    border: 1px solid #E2E8F0;
    border-left: 3px solid #1E293B;
    border-radius: 4px;
    padding: 10px 14px;
    font-size: 13.5px;
    color: #1E293B;
    margin: 10px 0 14px;
}
.bc-file span { color: #64748B; margin-left: 8px; }

/* ボタン */
.stButton > button, .stDownloadButton > button {
    border-radius: 4px;
    font-weight: 500;
    letter-spacing: 0.04em;
    padding: 0.6rem 1rem;
}

/* 進行表示 */
.bc-progress {
    background: #FFFFFF;
    border: 1px solid #E2E8F0;
    border-radius: 6px;
    padding: 20px 22px 18px;
    margin: 14px 0;
}
.bc-steps { display: flex; align-items: center; gap: 0; }
.bc-step { display: flex; align-items: center; gap: 8px; font-size: 13px; color: #94A3B8; white-space: nowrap; }
.bc-dot {
    width: 24px; height: 24px; border-radius: 50%;
    display: inline-flex; align-items: center; justify-content: center;
    font-size: 12px; font-weight: 700;
    border: 1.5px solid #CBD5E1; color: #94A3B8; background: #FFFFFF;
}
.bc-step.done { color: #1E293B; }
.bc-step.done .bc-dot { background: #1E293B; border-color: #1E293B; color: #FFFFFF; }
.bc-step.active { color: #1E293B; font-weight: 700; }
.bc-step.active .bc-dot { border-color: #B08D57; color: #B08D57; border-width: 2px; }
.bc-line { flex: 1; height: 1px; background: #E2E8F0; margin: 0 12px; min-width: 16px; }
.bc-line.done { background: #1E293B; }
.bc-msg { margin-top: 14px; font-size: 13px; color: #475569; line-height: 1.6; }
.bc-msg.wait { color: #92400E; }

/* 結果 */
.bc-result-head {
    display: flex; justify-content: space-between; align-items: baseline;
    border-bottom: 2px solid #1E293B;
    padding-bottom: 8px; margin: 34px 0 14px;
}
.bc-rh-title { margin: 0; font-size: 18px; color: #1E293B; font-weight: 700; }
.bc-result-head span { font-size: 12px; color: #64748B; }

.stTabs [data-baseweb="tab-list"] { gap: 4px; border-bottom: 1px solid #E2E8F0; }
.stTabs [data-baseweb="tab"] { padding: 8px 18px; }
.stTabs [data-baseweb="tab"] p { font-size: 14px !important; font-weight: 500; }
.stTabs [role="tabpanel"] {
    background: #FFFFFF;
    border: 1px solid #E2E8F0;
    border-top: none;
    border-radius: 0 0 6px 6px;
    padding: 22px 28px;
}
.stTabs [role="tabpanel"] p, .stTabs [role="tabpanel"] li {
    font-size: 14.5px; line-height: 1.85; color: #1E293B;
}
.stTabs [role="tabpanel"] h3, .stTabs [role="tabpanel"] h4 {
    color: #1E293B; border-left: 3px solid #B08D57; padding: 2px 0 2px 12px; margin: 0.4em 0 0.8em;
    font-size: 19px !important; font-weight: 700;
}
.stTabs [role="tabpanel"] strong { color: #0F172A; }
.stTabs [role="tabpanel"] table { font-size: 13.5px; }

/* 管理画面 */
.bc-admin {
    font-size: 12px; font-weight: 700; color: #B08D57; letter-spacing: 0.16em;
    margin-bottom: 4px;
}
</style>
""", unsafe_allow_html=True)


# =========================================================
# 進行表示
# =========================================================
STEPS = ["音声を転送", "音声を準備", "Geminiで解析"]


def render_progress(ph, step, message, waiting=False):
    """step: 0〜2 = 実行中の段, 3 = 全段完了"""
    parts = []
    for i, name in enumerate(STEPS):
        if i < step:
            cls, mark = "done", "✓"
        elif i == step:
            cls, mark = "active", str(i + 1)
        else:
            cls, mark = "", str(i + 1)
        parts.append(f'<div class="bc-step {cls}"><span class="bc-dot">{mark}</span>{name}</div>')
        if i < len(STEPS) - 1:
            parts.append(f'<div class="bc-line {"done" if i < step else ""}"></div>')
    msg_cls = "bc-msg wait" if waiting else "bc-msg"
    ph.markdown(
        f'<div class="bc-progress"><div class="bc-steps">{"".join(parts)}</div>'
        f'<div class="{msg_cls}">{message}</div></div>',
        unsafe_allow_html=True,
    )


def wait_with_message(ph, step, seconds, reason):
    """待ち時間を残り秒数つきで表示しながら待つ"""
    end = time.time() + seconds
    while True:
        remain = int(end - time.time())
        if remain <= 0:
            break
        render_progress(ph, step, f"{reason}。<b>{remain}秒後に再試行します。</b>そのままお待ちください。", waiting=True)
        time.sleep(1)


# =========================================================
# Gemini呼び出し（503対策）
# =========================================================
def upload_with_retry(client, path, mime_type, ph):
    """音声アップロード自体も混雑で失敗することがあるので再試行する"""
    last_error = None
    for attempt in range(1, 4):
        try:
            render_progress(ph, 0, f"音声ファイルを転送しています…（{attempt}/3回目）")
            return client.files.upload(file=path, config={"mime_type": mime_type})
        except genai_errors.APIError as e:
            last_error = e
            if e.code not in RETRYABLE_CODES or attempt == 3:
                raise
            wait_with_message(ph, 0, BASE_WAIT_SEC * attempt, f"転送が混雑しています（エラー{e.code}）")
    raise last_error


def wait_until_active(client, file_ref, ph):
    """サーバー側の音声処理完了を待つ。上限時間とFAILEDを検知する"""
    start = time.time()
    while file_ref.state.name == "PROCESSING":
        elapsed = int(time.time() - start)
        if elapsed > FILE_PROCESS_TIMEOUT_SEC:
            raise TimeoutError("音声ファイルの処理が時間内に終わりませんでした。")
        render_progress(ph, 1, f"Gemini側で音声を準備しています…（{elapsed}秒経過）")
        time.sleep(5)
        file_ref = client.files.get(name=file_ref.name)
    if file_ref.state.name != "ACTIVE":
        raise RuntimeError(f"音声ファイルの処理に失敗しました（状態: {file_ref.state.name}）")
    return file_ref


def generate_with_retry(client, contents, ph):
    """候補モデルを順に試し、混雑時は指数バックオフで再試行する"""
    last_error = None
    for model in MODEL_CANDIDATES:
        for attempt in range(1, MAX_ATTEMPTS_PER_MODEL + 1):
            try:
                render_progress(ph, 2, f"{model} で解析しています…（{attempt}/{MAX_ATTEMPTS_PER_MODEL}回目）。1時間の音声で1〜3分ほどかかります。")
                response = client.models.generate_content(model=model, contents=contents)
                if not response.text:
                    raise RuntimeError("Geminiから空の応答が返りました。")
                return response.text, model
            except genai_errors.APIError as e:
                last_error = e
                if e.code == 404:
                    break  # このモデル名は存在しない → 待たずに次の候補へ
                if e.code not in RETRYABLE_CODES:
                    raise  # 再試行しても直らないエラー
                if attempt == MAX_ATTEMPTS_PER_MODEL:
                    break  # このモデルは諦めて次の候補へ
                wait = min(MAX_WAIT_SEC, BASE_WAIT_SEC * (2 ** (attempt - 1))) + random.uniform(0, 3)
                wait_with_message(ph, 2, wait, f"{model} が混雑しています（エラー{e.code}）")
    if last_error is not None:
        raise last_error
    raise RuntimeError("利用可能なモデルが見つかりませんでした。")


# =========================================================
# 出力の整形（画面表示・Word）
# =========================================================
SECTION_RE = re.compile(r'^[\s#>*]*■\s*([1-4])\s*[\.．、]', re.M)
SECTION_NAMES = {"1": "議事録", "2": "Gemini N°1", "3": "Gemini N°2", "4": "質問"}


def split_sections(text):
    """「■ 1.」〜「■ 4.」の見出しで分割する。見つからなければ全文1枚"""
    matches = list(SECTION_RE.finditer(text))
    if len(matches) < 2:
        return [("全文", text)]
    sections = []
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        sections.append((SECTION_NAMES.get(m.group(1), f"項目{m.group(1)}"), text[m.start():end].strip()))
    sections.append(("全文", text))
    return sections


def to_display_md(body):
    """画面表示用：■見出しを見出し化し、単独改行が潰れないようにする"""
    out = []
    for line in body.splitlines():
        s = line.strip()
        if re.match(r'^[#*\s]*■', s):
            out.append("### " + re.sub(r'^[#*\s]*', '', s).replace("**", ""))
        elif s == "":
            out.append("")
        elif s.startswith("|"):
            out.append(line)
        else:
            out.append(line.rstrip() + "  ")
    return "\n".join(out)


def _set_font(style, size=None, bold=None, color=None):
    style.font.name = "Yu Gothic"
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), "游ゴシック")
    rfonts.set(qn('w:ascii'), "Yu Gothic")
    rfonts.set(qn('w:hAnsi'), "Yu Gothic")
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = color


def _add_runs(paragraph, text):
    """**太字** を反映して文字を入れる"""
    for i, chunk in enumerate(re.split(r'\*\*', text)):
        if chunk:
            run = paragraph.add_run(chunk.replace("*", "") if i % 2 == 0 else chunk)
            run.bold = (i % 2 == 1)


def build_docx(text, title, meta):
    doc = Document()
    navy = RGBColor(0x1E, 0x29, 0x3B)
    _set_font(doc.styles["Normal"], size=10.5)
    _set_font(doc.styles["Title"], size=20, bold=True, color=navy)
    for lv, size in ((1, 15), (2, 13), (3, 11.5)):
        _set_font(doc.styles[f"Heading {lv}"], size=size, bold=True, color=navy)
    for name in ("List Bullet", "List Bullet 2"):
        _set_font(doc.styles[name], size=10.5)

    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph()
    r = p.add_run(meta)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        s = raw.strip()
        i += 1
        if not s or re.fullmatch(r'[-*_]{3,}', s):
            continue
        # 表（| a | b |）
        if s.startswith("|"):
            rows = [s]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows
                     if not re.fullmatch(r'\|?[\s:\-|]+\|?', r)]
            if cells:
                ncol = max(len(c) for c in cells)
                table = doc.add_table(rows=len(cells), cols=ncol)
                table.style = "Table Grid"
                for ri, row in enumerate(cells):
                    for ci in range(ncol):
                        cell_p = table.cell(ri, ci).paragraphs[0]
                        _add_runs(cell_p, row[ci] if ci < len(row) else "")
                        if ri == 0:
                            for run in cell_p.runs:
                                run.bold = True
            continue
        m = re.match(r'^(#{1,6})\s+(.*)', s)
        if m:
            body = m.group(2).replace("**", "")
            level = 1 if body.startswith("■") else min(len(m.group(1)), 3)
            doc.add_heading(body, level=level)
            continue
        if re.match(r'^\**■', s):
            doc.add_heading(s.replace("**", ""), level=1)
            continue
        indent = len(raw) - len(raw.lstrip())
        bm = re.match(r'^(?:[-*•]\s+|・)(.*)', s)
        if bm:
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            _add_runs(doc.add_paragraph(style=style), bm.group(1))
            continue
        _add_runs(doc.add_paragraph(), s)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# =========================================================
# 画面
# =========================================================
query_params = st.query_params
is_admin = False
if "admin" in query_params:
    param_val = query_params["admin"]
    if isinstance(param_val, list):
        is_admin = len(param_val) > 0 and param_val == "true"
    else:
        is_admin = param_val == "true"

if "title_val" not in st.session_state:
    st.session_state["title_val"] = "Kobunsha BIZ-CREA"
if "desc_val" not in st.session_state:
    st.session_state["desc_val"] = "会議音声から、議事録と新事業の企画案ドラフトを生成します。1時間を超える音声にも対応しています。"
if "pre_val" not in st.session_state:
    st.session_state["pre_val"] = "・現在の出版業界は非常に厳しい。既存の「紙」や、失敗に終わった消耗戦のコタツ記事デジタルから完全に脱却する。\n・紙、デジタルにとらわれず、イベント、スクールビジネス、終活ビジネス等々の「光文社が手がけなかった」新規事業を生み出す。\n・既存雑誌・書籍と一切関係ないK-POP関係の出版物、グッズ販売、イベント等を行ってきた。\n・フィジカルか、AIか、動画か、SNSかを問わず、新たなビジネスイメージを作りたい。\n・もちろん、書籍やIPといった出版社の事業も大切。\n・ただし「出版社がやるビジネス」のイメージを守る。ラーメン屋やキャバクラを経営してどうする？\n・海外の最新事例の構造も参照すること。"
if "target_val" not in st.session_state:
    st.session_state["target_val"] = "・年々高齢化する40～60代女性という光文社女性誌の既存読者だけでなく、もっと若い世代、男性、オタク層、推し活層、在日外国人その他の「新たなマーケット」を開拓する。"
if "ng_val" not in st.session_state:
    st.session_state["ng_val"] = "・コンサル的な「お行儀の良い一般論」「誰でも言える綺麗事（絆、おもてなし）」。\n・予算を無駄に食う割に、現場を動かせないハコモノ的な大裟な提案。"

def sync_title(): st.session_state["title_val"] = st.session_state["k1"]
def sync_desc(): st.session_state["desc_val"] = st.session_state["k2"]
def sync_pre(): st.session_state["pre_val"] = st.session_state["k3"]
def sync_target(): st.session_state["target_val"] = st.session_state["k4"]
def sync_ng(): st.session_state["ng_val"] = st.session_state["k5"]

if is_admin:
    left_col, right_col = st.columns(2, gap="large")
else:
    _, right_col, _ = st.columns([1, 2.2, 1])

if is_admin:
    with left_col:
        st.markdown('<div class="bc-admin">ADMIN</div>', unsafe_allow_html=True)
        st.markdown("#### サイト文言・前提条件の編集")
        st.text_input("システムのメインタイトル", value=st.session_state["title_val"], key="k1", on_change=sync_title)
        st.text_area("システムの説明文", value=st.session_state["desc_val"], key="k2", height=70, on_change=sync_desc)
        st.markdown('<div class="bc-label">戦略文脈（プロンプト）</div>', unsafe_allow_html=True)
        st.text_area("前提条件", value=st.session_state["pre_val"], key="k3", height=220, on_change=sync_pre)
        st.text_area("狙うべきターゲット層", value=st.session_state["target_val"], key="k4", height=100, on_change=sync_target)
        st.text_area("排除すべきNG項目", value=st.session_state["ng_val"], key="k5", height=100, on_change=sync_ng)

with right_col:
    st.markdown(f"""
    <div class="bc-header">
        <p class="bc-kicker">KOBUNSHA BUSINESS CREATION</p>
        <div class="bc-title">{html.escape(st.session_state["title_val"])}</div>
        <p class="bc-desc">{html.escape(st.session_state["desc_val"])}</p>
    </div>
    """, unsafe_allow_html=True)

    st.markdown('<div class="bc-label">会議の録音ファイル（MP3 / M4A / WAV）</div>', unsafe_allow_html=True)
    uploaded_file = st.file_uploader("会議の録音ファイル", type=["mp3", "m4a", "wav"], label_visibility="collapsed")

    if uploaded_file:
        st.markdown(
            f'<div class="bc-file">{html.escape(uploaded_file.name)}<span>{uploaded_file.size/1024/1024:.1f} MB</span></div>',
            unsafe_allow_html=True,
        )

        if st.button("議事録と企画案を生成する", type="primary", use_container_width=True):
            st.session_state.pop("result", None)
            progress_ph = st.empty()

            api_key = os.environ.get("GEMINI_API_KEY")
            if not api_key:
                st.error("認証エラー: GEMINI_API_KEYが登録されていません。")
                st.stop()

            client = genai.Client(api_key=api_key)
            temp_path = None
            audio_file_ref = None

            try:
                ext = os.path.splitext(uploaded_file.name)[1].lower()

                # 利用者ごとに別名の一時ファイルを作る（同時利用での上書き防止）
                with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                    tmp.write(uploaded_file.getbuffer())
                    temp_path = tmp.name

                mime_map = {'.wav': 'audio/wav', '.mp3': 'audio/mp3', '.m4a': 'audio/m4a'}
                mime_type = mime_map.get(ext, 'audio/mp3')

                # 音声は1回だけアップロードし、再試行ではこれを使い回す
                audio_file_ref = upload_with_retry(client, temp_path, mime_type, progress_ph)
                audio_file_ref = wait_until_active(client, audio_file_ref, progress_ph)

                prompt = f"""
                あなたは光文社ビジネスクリエーション部の会議に同席し、検討のための「素材」を提供する役割です。
                最終的な判断は、すべて人間が行います。あなたは結論・評価・判断を下さず、人間が考えるための材料を出すことに徹してください。
                添付された会議音声について、【議事録】と【新事業の企画案（素材）】を出力してください。
                企画案は、性格の異なる2人の企画者「Gemini N°1」「Gemini N°2」が2案ずつ、計4案を出し、最後に互いの案へ質問をぶつけること。
                本日は{datetime.now(JST):%Y年%m月%d日}です。

                【全体のルール】
                ・出力は日本語のみ。英字は固有名詞・ブランド名・事例名に限る。簡体字・繁体字・ハングル等の外国語の文字は使わない。
                ・出力は「■ 1. 議事録」から書き始めること。前置き・挨拶・締めの文は書かない。
                ・金額・人数などの数字は、万・億・兆の単位を使って読みやすく書くこと（例：14億4,000万円、3,000人、3兆円）。0を並べた表記はしない。

                【議事録（■1）のルール】
                ・議事録は、音声で実際に話されたことの記録に限る。あなたの解釈・評価・提案・結論は一切書かない。
                ・話された内容を、会議の方向性や下記の前提条件に合わせて言い換えたり、格上げしたりしない。
                ・音声で言及されていない数字・日付・期限・固有名詞は書かない。
                ・議事録の中で、後に続くGemini N°1・N°2・質問に言及しない。
                ・下記の「前提条件」「ターゲット層」「企画におけるNG」は企画案（■2〜■4）だけに適用し、議事録には適用しない。

                【企画案（■2〜■4）のルール】
                ・4案はすべて、この会議で話された議題・発言を起点に発想すること。会議で話されていない領域の案は出さない。
                ・前提条件・ターゲット層・NGは、会議を起点にした発想を光文社の事業としてどう形にするかの方向づけにだけ使う。前提条件に例として挙がっている領域（K-POP、終活など）を、会議の内容と無関係に持ち込まない。
                ・4案のテーマ・事業領域を互いに重複させない。
                ・「参照すべき海外最新事例」には、実在することに確信がある企業・サービス・事例だけを書く。事例名・企業名・サービス名を創作することは絶対に禁止。少しでも実在に確信が持てない場合は、何も書かずに「確実な該当事例なし」と書くこと。もっともらしい架空の事例を書くより、空欄のほうが価値がある。

                ■ 前提条件（企画案のみに適用）：
                {st.session_state["pre_val"]}

                ■ ターゲット層（企画案のみに適用）：
                {st.session_state["target_val"]}

                ■ 企画におけるNG（企画案のみに適用）：
                {st.session_state["ng_val"]}

                以下のフォーマットを厳格に守って出力すること：

                ■ 1. 議事録
                ・議題ごとの発言要旨：話された順に、議題ごとに要旨をまとめる。話者は「話者A」「話者B」等の記号だけで示す。人名・肩書・所属は、音声で名乗られていても一切書かない（聞き取り違いによる誤記を防ぐため）。
                ・決定事項：参加者が明確に合意した発言がある場合のみ書く。なければ「決定事項なし」と書く。
                ・争点：意見が分かれた点について、問いの形だけで済ませず、どの話者が何を主張したかを書く。なければ「なし」と書く。
                　争点は、異なる話者が互いに異なる意見を実際に述べた場合のみ書く。対立を推測で作らない。
                　一人の話者が「〜は必要だが、〜という問題もある」のように両面を述べた場合は、それを別々の話者の対立として書かない。発言要旨の中で、その話者の発言として両面をそのまま書く。
                ・Next Actions：音声で担当・期限・作業が言及されたもののみ書く。なければ「言及なし」と書く。

                ■ 2. Gemini N°1
                N°1は、論理と数字に徹底的に振り切った企画者である。この視点で2案を出す。
                各案：タイトル／発想の起点（議事録のどの議題・発言から発想したかを1行で）／核心アイデア／数値の根拠（市場規模・想定顧客数・収支の目安）／参照すべき海外最新事例／具体的な実行ステップ
                ※数値はすべて「推定」と明記し、各数値の直後に計算式（前提となる数値×割合＝結果）で根拠を示すこと。根拠を示せない数値は書かないこと。
                ※統計名・調査会社名・官公庁名などの出典名は書かないこと。前提となる数値そのものも推定であることを明記すること。
                ※計算結果は必ず検算し、桁区切りとパーセントを正しく書くこと。営業利益は金額で、利益率は売上高に対する割合（100%未満）で書くこと。

                ■ 3. Gemini N°2
                N°2は、人が集まり、思わず誰かに話したくなる面白さに徹底的に振り切った企画者である。この視点で2案を出す。
                各案：タイトル／発想の起点（議事録のどの議題・発言から発想したかを1行で）／核心アイデア／話題化の仕掛け（誰が、なぜ人に話したくなるのか）／参照すべき海外最新事例／具体的な実行ステップ

                ※見出しには「Gemini N°1」「Gemini N°2」とだけ書くこと。N°1とN°2の性格や視点を説明する文は、出力のどこにも書かないこと。読み手が案の中身から2人の違いに気づくようにする。

                ■ 4. 質問
                ・N°1からN°2へ：N°2の各案に対し、N°1の視点から、相手が答えに詰まる急所を突く質問を2つずつ
                ・N°2からN°1へ：N°1の各案に対し、N°2の視点から、相手が答えに詰まる急所を突く質問を2つずつ
                ※単なる確認や情報の追加要求（「ターゲットを詳しく教えてください」等）は禁止。案の前提や成立条件そのものを揺さぶる質問にすること。
                ※質問への回答、改善案、まとめは一切書かないこと。質問だけで終えること。
                """

                output_text, used_model = generate_with_retry(client, [audio_file_ref, prompt], progress_ph)

                now = datetime.now(JST)
                meta = f"生成日時 {now:%Y年%m月%d日 %H:%M}　／　音声 {uploaded_file.name}　／　モデル {used_model}"
                st.session_state["result"] = {
                    "text": output_text,
                    "model": used_model,
                    "meta": meta,
                    "base": f"光文社ビジクリ部_議事録_企画案_{now:%Y%m%d_%H%M}",
                    "docx": build_docx(output_text, "議事録・新事業企画案", meta),
                }
                render_progress(progress_ph, 3, "完了しました。")

            except genai_errors.APIError as e:
                progress_ph.empty()
                if e.code in RETRYABLE_CODES:
                    st.error(f"Geminiの混雑が続いており、全モデルで再試行しましたが完了できませんでした（エラー{e.code}）。時間をおいて再実行してください。")
                else:
                    st.error(f"Geminiでエラーが発生しました（エラー{e.code}）: {e}")
            except Exception as e:
                progress_ph.empty()
                st.error(f"エラーが発生しました: {e}")
            finally:
                # 成功・失敗にかかわらず、Gemini側とサーバー側の音声を必ず削除
                if audio_file_ref is not None:
                    try:
                        client.files.delete(name=audio_file_ref.name)
                    except Exception:
                        pass
                if temp_path and os.path.exists(temp_path):
                    os.remove(temp_path)

    # 結果はセッションに保持する（ダウンロードボタンを押しても消えない）
    result = st.session_state.get("result")
    if result:
        st.markdown(
            f'<div class="bc-result-head"><div class="bc-rh-title">生成結果</div><span>{html.escape(result["meta"])}</span></div>',
            unsafe_allow_html=True,
        )
        c1, c2 = st.columns(2)
        with c1:
            st.download_button(
                "TXTでダウンロード",
                data=result["text"],
                file_name=result["base"] + ".txt",
                mime="text/plain",
                use_container_width=True,
                key="dl_txt",
            )
        with c2:
            st.download_button(
                "Wordでダウンロード",
                data=result["docx"],
                file_name=result["base"] + ".docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
                use_container_width=True,
                key="dl_docx",
            )

        sections = split_sections(result["text"])
        tabs = st.tabs([name for name, _ in sections])
        for tab, (name, body) in zip(tabs, sections):
            with tab:
                st.markdown(to_display_md(body))
